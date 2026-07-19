"""Render ATS-friendly tailored-resume Markdown as DOCX."""

from __future__ import annotations

from io import BytesIO
from html import escape
import re


def render_resume_docx(markdown: str) -> bytes:
    try:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("python-docx is required to generate resume files.") from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, before, after, color in (
        ("Heading 1", 16, 18, 10, "2E74B5"),
        ("Heading 2", 13, 14, 7, "2E74B5"),
        ("Heading 3", 12, 10, 5, "1F4D78"),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Resume Header" not in styles:
        header_style = styles.add_style("Resume Header", WD_STYLE_TYPE.PARAGRAPH)
    else:
        header_style = styles["Resume Header"]
    header_style.font.name = "Calibri"
    header_style.font.size = Pt(18)
    header_style.font.bold = True
    header_style.font.color.rgb = RGBColor.from_string("0B2545")
    header_style.paragraph_format.space_after = Pt(4)
    header_style.paragraph_format.keep_with_next = True

    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Resume Header")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.add_run(_plain_text(line[2:]))
        elif line.startswith("## "):
            paragraph = document.add_paragraph(_plain_text(line[3:]), style="Heading 2")
            _add_bottom_border(paragraph, OxmlElement, qn)
        elif line.startswith("### "):
            document.add_paragraph(_plain_text(line[4:]), style="Heading 3")
        elif re.match(r"^[-*]\s+", line):
            bullet_text = _plain_text(re.sub(r"^[-*]\s+", "", line))
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.375)
            paragraph.paragraph_format.first_line_indent = Inches(-0.188)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.25
            paragraph.add_run(f" {bullet_text}")
        else:
            document.add_paragraph(_plain_text(line))

    properties = document.core_properties
    properties.title = "Tailored Resume"
    properties.subject = "JobPilot generated resume"
    properties.author = "JobPilot"
    properties.keywords = "resume"
    properties.comments = "Generated from user-provided evidence for human review."

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_resume_in_original_docx(original_file: bytes, markdown: str) -> bytes:
    """Rewrite body text in a copy of a DOCX while retaining its document package and layout."""
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise RuntimeError("python-docx is required to preserve DOCX resume templates.") from exc

    blocks = _markdown_blocks(markdown)
    if not blocks:
        raise ValueError("Cannot render an empty resume.")
    try:
        document = Document(BytesIO(original_file))
    except Exception as exc:
        raise ValueError("The original DOCX template could not be opened.") from exc

    header_footer_text = " ".join(
        paragraph.text
        for section in document.sections
        for area in (section.header, section.footer)
        for paragraph in area.paragraphs
        if paragraph.text.strip()
    ).casefold()
    if blocks and blocks[0][0] == "title" and header_footer_text:
        title_tokens = [token for token in re.findall(r"[a-z0-9]+", blocks[0][1].casefold()) if len(token) > 1]
        if title_tokens and all(token in header_footer_text for token in title_tokens):
            blocks = blocks[1:]

    slots = _document_body_paragraphs(document, Paragraph, Table)
    editable_slots = [paragraph for paragraph in slots if paragraph.text.strip()]
    if not editable_slots:
        raise ValueError("The original DOCX does not contain editable body text.")

    for paragraph, (_, text) in zip(editable_slots, blocks):
        _replace_paragraph_text_preserving_layout(paragraph, text)

    used_count = min(len(editable_slots), len(blocks))
    for paragraph in editable_slots[used_count:]:
        _remove_paragraph_text_preserving_objects(paragraph)

    if len(blocks) > len(editable_slots):
        for kind, text in blocks[len(editable_slots):]:
            paragraph = document.add_paragraph()
            if kind == "bullet":
                paragraph.style = _available_style(document, "List Bullet", "Normal")
            elif kind in {"section", "entry"}:
                paragraph.style = _available_style(
                    document,
                    "Heading 2" if kind == "section" else "Heading 3",
                    "Normal",
                )
            paragraph.add_run(text)

    properties = document.core_properties
    properties.title = "Tailored Resume"
    properties.subject = "JobPilot tailored resume using the candidate's original DOCX layout"
    properties.comments = (
        "Content was tailored from user-provided evidence. Original DOCX layout was preserved on a best-effort basis."
    )
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_resume_pdf(markdown: str) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise RuntimeError("reportlab is required to generate PDF resume files.") from exc

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title="Tailored Resume",
        author="JobPilot",
        subject="JobPilot generated resume",
    )
    base = getSampleStyleSheet()
    styles = {
        "title": ParagraphStyle(
            "ResumeTitle",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=21,
            textColor=colors.HexColor("#0B2545"),
            spaceAfter=5,
            alignment=TA_LEFT,
        ),
        "section": ParagraphStyle(
            "ResumeSection",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=14,
            textColor=colors.HexColor("#2E74B5"),
            spaceBefore=12,
            spaceAfter=4,
            keepWithNext=True,
        ),
        "entry": ParagraphStyle(
            "ResumeEntry",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=10.5,
            leading=13,
            textColor=colors.HexColor("#1F4D78"),
            spaceBefore=7,
            spaceAfter=3,
            keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "ResumeBody",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=10.5,
            leading=13.5,
            textColor=colors.black,
            spaceAfter=5,
        ),
        "bullet": ParagraphStyle(
            "ResumeBullet",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=10.5,
            leading=13.5,
            textColor=colors.black,
            leftIndent=14,
            firstLineIndent=-9,
            spaceAfter=3,
        ),
    }

    story = []
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            story.append(Paragraph(_pdf_text(line[2:]), styles["title"]))
        elif line.startswith("## "):
            story.append(Spacer(1, 2))
            story.append(Paragraph(_pdf_text(line[3:]), styles["section"]))
            story.append(
                HRFlowable(
                    width="100%",
                    thickness=0.6,
                    color=colors.HexColor("#D9E2F3"),
                    spaceBefore=0,
                    spaceAfter=5,
                )
            )
        elif line.startswith("### "):
            story.append(Paragraph(_pdf_text(line[4:]), styles["entry"]))
        elif re.match(r"^[-*]\s+", line):
            bullet_text = _pdf_text(re.sub(r"^[-*]\s+", "", line))
            story.append(Paragraph(f"-&nbsp;&nbsp;{bullet_text}", styles["bullet"]))
        else:
            story.append(Paragraph(_pdf_text(line), styles["body"]))

    if not story:
        raise ValueError("Cannot render an empty resume.")
    document.build(story)
    return output.getvalue()


def _plain_text(value: str) -> str:
    value = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", value)
    value = re.sub(r"(\*\*|__|\*|`)", "", value)
    return re.sub(r"\s+", " ", value).strip()


def _markdown_blocks(markdown: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            blocks.append(("title", _plain_text(line[2:])))
        elif line.startswith("## "):
            blocks.append(("section", _plain_text(line[3:])))
        elif line.startswith("### "):
            blocks.append(("entry", _plain_text(line[4:])))
        elif re.match(r"^[-*]\s+", line):
            blocks.append(("bullet", _plain_text(re.sub(r"^[-*]\s+", "", line))))
        else:
            blocks.append(("body", _plain_text(line)))
    return [(kind, text) for kind, text in blocks if text]


def _document_body_paragraphs(document, paragraph_type, table_type) -> list:
    paragraphs = []
    seen_cells: set[int] = set()
    for block in document.iter_inner_content():
        if isinstance(block, paragraph_type):
            paragraphs.append(block)
        elif isinstance(block, table_type):
            for row in block.rows:
                for cell in row.cells:
                    cell_id = id(cell._tc)
                    if cell_id in seen_cells:
                        continue
                    seen_cells.add(cell_id)
                    paragraphs.extend(cell.paragraphs)
    return paragraphs


def _replace_paragraph_text_preserving_layout(paragraph, text: str) -> None:
    text_nodes = paragraph._p.xpath(".//w:t")
    if text_nodes:
        text_nodes[0].text = text
        for node in text_nodes[1:]:
            node.text = ""
        return
    paragraph.add_run(text)


def _remove_paragraph_text_preserving_objects(paragraph) -> None:
    for node in paragraph._p.xpath(".//w:t"):
        node.text = ""


def _available_style(document, preferred: str, fallback: str) -> str:
    return preferred if preferred in document.styles else fallback


def _pdf_text(value: str) -> str:
    text = _plain_text(value)
    text = (
        text.replace("\u2010", "-")
        .replace("\u2011", "-")
        .replace("\u2012", "-")
        .replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u2212", "-")
    )
    return escape(text)


def _add_bottom_border(paragraph, oxml_element, qn) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = oxml_element("w:pBdr")
        properties.append(borders)
    bottom = oxml_element("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "D9E2F3")
    borders.append(bottom)
