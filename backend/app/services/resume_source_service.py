"""Validate, extract, and persist original resume uploads."""

from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
from zipfile import BadZipFile, ZipFile

from sqlalchemy.orm import Session

from app.models.resume import UserResume
from app.repositories.resume_repository import upsert_user_resume_file
from app.storage.resume_storage import load_resume_file, prepare_resume_file


DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_CONTENT_TYPE = "application/pdf"
ALLOWED_CONTENT_TYPES = {DOCX_CONTENT_TYPE, PDF_CONTENT_TYPE}
MAX_RESUME_FILE_SIZE = 10 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_SIZE = 60 * 1024 * 1024
MAX_PDF_PAGES = 20


class ResumeSourceError(ValueError):
    """Raised when an original resume upload cannot be safely processed."""


def save_original_resume(
    db: Session,
    *,
    user_id: int,
    filename: str,
    content_type: str | None,
    file_data: bytes,
) -> UserResume:
    safe_filename, normalized_type = validate_original_resume(
        filename=filename,
        content_type=content_type,
        file_data=file_data,
    )
    resume_text = extract_resume_text(file_data, normalized_type)
    if len(resume_text) < 20:
        raise ResumeSourceError(
            "The uploaded file does not contain enough selectable text. "
            "Scanned PDFs need OCR before they can be tailored."
        )
    stored = prepare_resume_file(
        user_id=user_id,
        filename=safe_filename,
        data=file_data,
        category="source",
    )
    return upsert_user_resume_file(
        db,
        user_id=user_id,
        resume_text=resume_text,
        filename=safe_filename,
        content_type=normalized_type,
        file_data=file_data,
        file_sha256=sha256(file_data).hexdigest(),
        storage_provider=stored.provider,
        storage_key=stored.key,
    )


def validate_original_resume(
    *,
    filename: str,
    content_type: str | None,
    file_data: bytes,
) -> tuple[str, str]:
    if not file_data:
        raise ResumeSourceError("The uploaded resume is empty.")
    if len(file_data) > MAX_RESUME_FILE_SIZE:
        raise ResumeSourceError("Resume files must be 10 MB or smaller.")

    original_name = Path(filename or "resume").name
    safe_filename = re.sub(r"[^A-Za-z0-9._ -]+", "_", original_name).strip(" .")[:255]
    if not safe_filename:
        raise ResumeSourceError("The uploaded resume filename is invalid.")
    suffix = Path(safe_filename).suffix.casefold()
    if file_data.startswith(b"%PDF-"):
        detected_type = PDF_CONTENT_TYPE
        expected_suffix = ".pdf"
    elif file_data.startswith(b"PK") and suffix == ".docx":
        detected_type = DOCX_CONTENT_TYPE
        expected_suffix = ".docx"
        _validate_docx_archive(file_data)
    else:
        raise ResumeSourceError("Only valid DOCX and text-based PDF resumes are supported.")

    if suffix != expected_suffix:
        raise ResumeSourceError(f"The file contents do not match the {expected_suffix} extension.")
    if content_type and content_type not in ALLOWED_CONTENT_TYPES and content_type != "application/octet-stream":
        raise ResumeSourceError("The uploaded resume content type is not supported.")
    return safe_filename, detected_type


def extract_resume_text(file_data: bytes, content_type: str) -> str:
    if content_type == DOCX_CONTENT_TYPE:
        return _extract_docx_text(file_data)
    if content_type == PDF_CONTENT_TYPE:
        return _extract_pdf_text(file_data)
    raise ResumeSourceError("Unsupported resume content type.")


def load_original_resume_file(resume: UserResume) -> bytes:
    if not resume.original_storage_provider:
        raise ResumeSourceError("The original resume file is not available.")
    return load_resume_file(resume)


def _extract_docx_text(file_data: bytes) -> str:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise RuntimeError("python-docx is required to read DOCX resumes.") from exc
    try:
        document = Document(BytesIO(file_data))
    except Exception as exc:
        raise ResumeSourceError("The DOCX file could not be opened.") from exc

    lines: list[str] = []
    seen_cells: set[int] = set()
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            _append_text(lines, block.text)
        elif isinstance(block, Table):
            for row in block.rows:
                for cell in row.cells:
                    cell_id = id(cell._tc)
                    if cell_id in seen_cells:
                        continue
                    seen_cells.add(cell_id)
                    for paragraph in cell.paragraphs:
                        _append_text(lines, paragraph.text)
    for section in document.sections:
        for area in (section.header, section.footer):
            for paragraph in area.paragraphs:
                _append_text(lines, paragraph.text)
    return "\n".join(lines).strip()


def _extract_pdf_text(file_data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf is required to read PDF resumes.") from exc
    try:
        reader = PdfReader(BytesIO(file_data))
        if reader.is_encrypted:
            raise ResumeSourceError("Password-protected PDFs are not supported.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ResumeSourceError(f"PDF resumes cannot exceed {MAX_PDF_PAGES} pages.")
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except ResumeSourceError:
        raise
    except Exception as exc:
        raise ResumeSourceError("The PDF file could not be opened or read.") from exc
    return "\n\n".join(page for page in pages if page).strip()


def _validate_docx_archive(file_data: bytes) -> None:
    try:
        with ZipFile(BytesIO(file_data)) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                raise ResumeSourceError("The uploaded file is not a valid DOCX document.")
            uncompressed_size = sum(item.file_size for item in archive.infolist())
            if uncompressed_size > MAX_DOCX_UNCOMPRESSED_SIZE:
                raise ResumeSourceError("The DOCX expands beyond the allowed safety limit.")
    except BadZipFile as exc:
        raise ResumeSourceError("The uploaded DOCX archive is invalid.") from exc


def _append_text(lines: list[str], value: str) -> None:
    cleaned = re.sub(r"[ \t]+", " ", value).strip()
    if cleaned:
        lines.append(cleaned)
