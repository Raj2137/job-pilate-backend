from unittest import TestCase
from unittest.mock import patch

from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.database.session import Base
from app.models.job import Job
from app.models.user import User
from app.schemas.resume import TailoredResumeRequest, TailoredResumeResponse
from app.services.resume_artifact_service import (
    generate_and_save_tailored_resume,
    get_saved_tailored_resume,
    list_saved_tailored_resumes,
    tailored_resume_file,
)
from app.services.resume_document_service import render_resume_docx, render_resume_pdf
from app.services.resume_source_service import DOCX_CONTENT_TYPE, save_original_resume


class ResumeArtifactServiceTests(TestCase):
    def setUp(self) -> None:
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)
        self.user = User(email="raj@example.com", full_name="Raj Kumar", hashed_password="hashed")
        self.job = Job(
            source="greenhouse",
            external_id="artifact-job-1",
            title="Backend Engineer",
            company="Example Co",
            description="Build Python and FastAPI services.",
        )
        self.db.add_all([self.user, self.job])
        self.db.commit()
        self.db.refresh(self.user)
        self.db.refresh(self.job)

    def tearDown(self) -> None:
        self.db.close()
        self.engine.dispose()

    @patch("app.services.resume_artifact_service.render_resume_docx", return_value=b"docx-bytes")
    @patch("app.services.resume_artifact_service.tailor_resume_for_job")
    def test_generates_versions_stores_and_loads_docx(self, tailor, render) -> None:
        tailor.return_value = TailoredResumeResponse(
            job_id=self.job.id,
            job_title=self.job.title,
            company=self.job.company,
            source_resume_id=None,
            tailored_resume_markdown="# Raj Kumar\n## Experience\n- Built Python APIs.",
            headline="Backend Engineer",
            professional_summary="Backend engineer building Python APIs.",
            core_skills=["Python", "FastAPI"],
            keywords_incorporated=["Python"],
            unsupported_job_requirements=[],
            change_summary=["Prioritized Python experience."],
            truthfulness_warnings=[],
            estimated_alignment_score=82,
            provider="openai",
            model="gpt-test",
        )

        first = generate_and_save_tailored_resume(
            self.db,
            user=self.user,
            payload=TailoredResumeRequest(
                job_id=self.job.id,
                resume_text="Raj Kumar. Backend engineer building Python APIs.",
            ),
        )
        second = generate_and_save_tailored_resume(
            self.db,
            user=self.user,
            payload=TailoredResumeRequest(
                job_id=self.job.id,
                resume_text="Raj Kumar. Backend engineer building Python APIs.",
            ),
        )

        self.assertEqual(1, first.version)
        self.assertEqual(2, second.version)
        self.assertTrue(first.filename.endswith("-v1.docx"))
        self.assertEqual("database", first.storage_provider)
        self.assertEqual("ats", first.actual_render_mode)
        self.assertEqual("standardized", first.template_fidelity)
        self.assertEqual(len(b"docx-bytes"), first.file_size)
        record = get_saved_tailored_resume(self.db, user_id=self.user.id, resume_id=first.id)
        self.assertIsNotNone(record)
        self.assertEqual(b"docx-bytes", tailored_resume_file(record))
        self.assertEqual(2, len(list_saved_tailored_resumes(self.db, user_id=self.user.id, limit=10)))
        self.assertEqual(2, render.call_count)

    def test_renders_application_ready_docx(self) -> None:
        data = render_resume_docx(
            "# Raj Kumar\n"
            "Backend Engineer | Python | FastAPI\n\n"
            "## Experience\n"
            "### Backend Engineer — Example Co | 2024–Present\n"
            "- Built Python and FastAPI services.\n"
        )

        self.assertTrue(data.startswith(b"PK"))
        document = Document(__import__("io").BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Raj Kumar", text)
        self.assertIn("Built Python and FastAPI services.", text)

    @patch(
        "app.services.resume_artifact_service.render_resume_in_original_docx",
        return_value=b"preserved-docx",
    )
    @patch("app.services.resume_artifact_service.tailor_resume_for_job")
    def test_auto_mode_uses_saved_original_docx(self, tailor, render_original) -> None:
        source_document = Document()
        source_document.add_paragraph("Raj Kumar - Backend Engineer")
        source_document.add_paragraph("Built Python APIs for production systems.")
        source_output = __import__("io").BytesIO()
        source_document.save(source_output)
        saved = save_original_resume(
            self.db,
            user_id=self.user.id,
            filename="raj-resume.docx",
            content_type=DOCX_CONTENT_TYPE,
            file_data=source_output.getvalue(),
        )
        tailor.return_value = TailoredResumeResponse(
            job_id=self.job.id,
            job_title=self.job.title,
            company=self.job.company,
            source_resume_id=saved.id,
            tailored_resume_markdown="# Raj Kumar\n## Experience\n- Built Python APIs.",
            headline="Backend Engineer",
            professional_summary="Backend engineer building Python APIs.",
            core_skills=["Python"],
            keywords_incorporated=["Python"],
            unsupported_job_requirements=[],
            change_summary=["Prioritized Python experience."],
            truthfulness_warnings=[],
            estimated_alignment_score=82,
            provider="openai",
            model="gpt-test",
        )

        result = generate_and_save_tailored_resume(
            self.db,
            user=self.user,
            payload=TailoredResumeRequest(job_id=self.job.id),
        )

        self.assertEqual("original_docx", result.actual_render_mode)
        self.assertEqual("best_effort", result.template_fidelity)
        self.assertEqual("ats", result.pdf_render_mode)
        render_original.assert_called_once()

    def test_renders_application_ready_pdf(self) -> None:
        data = render_resume_pdf(
            "# Raj Kumar\n"
            "Backend Engineer | Python | FastAPI\n\n"
            "## Experience\n"
            "### Backend Engineer - Example Co | 2024-Present\n"
            "- Built Python and FastAPI services.\n"
        )

        self.assertTrue(data.startswith(b"%PDF-"))
        self.assertGreater(len(data), 1000)
