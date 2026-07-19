from io import BytesIO
from unittest import TestCase

from docx import Document
from reportlab.pdfgen.canvas import Canvas
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.database.session import Base
from app.models.user import User
from app.services.resume_document_service import render_resume_in_original_docx
from app.services.resume_source_service import (
    DOCX_CONTENT_TYPE,
    PDF_CONTENT_TYPE,
    ResumeSourceError,
    extract_resume_text,
    load_original_resume_file,
    save_original_resume,
)


def _docx_bytes() -> bytes:
    document = Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "Raj Kumar"
    document.add_heading("Experience", level=1)
    document.add_paragraph("Backend Engineer at Acme")
    document.add_paragraph("Built Python and FastAPI services.", style="List Bullet")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_bytes() -> bytes:
    output = BytesIO()
    canvas = Canvas(output)
    canvas.drawString(72, 760, "Raj Kumar - Backend Engineer")
    canvas.drawString(72, 740, "Built Python, FastAPI, and PostgreSQL services.")
    canvas.save()
    return output.getvalue()


class ResumeSourceServiceTests(TestCase):
    def setUp(self) -> None:
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)
        self.user = User(email="source@example.com", full_name="Raj Kumar", hashed_password="hashed")
        self.db.add(self.user)
        self.db.commit()
        self.db.refresh(self.user)

    def tearDown(self) -> None:
        self.db.close()
        self.engine.dispose()

    def test_saves_original_docx_and_reports_extracted_text(self) -> None:
        file_data = _docx_bytes()
        resume = save_original_resume(
            self.db,
            user_id=self.user.id,
            filename="raj-resume.docx",
            content_type=DOCX_CONTENT_TYPE,
            file_data=file_data,
        )

        self.assertIn("Built Python and FastAPI services.", resume.resume_text)
        self.assertEqual(DOCX_CONTENT_TYPE, resume.content_type)
        self.assertEqual(len(file_data), resume.original_file_size)
        self.assertEqual(64, len(resume.original_file_sha256))
        self.assertEqual(file_data, load_original_resume_file(resume))

    def test_extracts_selectable_pdf_text(self) -> None:
        extracted = extract_resume_text(_pdf_bytes(), PDF_CONTENT_TYPE)
        self.assertIn("Backend Engineer", extracted)
        self.assertIn("PostgreSQL", extracted)

    def test_rejects_spoofed_file_extension(self) -> None:
        with self.assertRaises(ResumeSourceError):
            save_original_resume(
                self.db,
                user_id=self.user.id,
                filename="resume.docx",
                content_type=DOCX_CONTENT_TYPE,
                file_data=_pdf_bytes(),
            )

    def test_tailored_docx_keeps_header_and_section_geometry(self) -> None:
        original = _docx_bytes()
        original_document = Document(BytesIO(original))
        rendered = render_resume_in_original_docx(
            original,
            "# Raj Kumar\n## Experience\n- Built Python APIs for production systems.",
        )
        tailored_document = Document(BytesIO(rendered))

        self.assertEqual(
            original_document.sections[0].top_margin,
            tailored_document.sections[0].top_margin,
        )
        self.assertEqual("Raj Kumar", tailored_document.sections[0].header.paragraphs[0].text)
        text = "\n".join(paragraph.text for paragraph in tailored_document.paragraphs)
        self.assertEqual("Experience", tailored_document.paragraphs[0].text)
        self.assertIn("Built Python APIs for production systems.", text)
